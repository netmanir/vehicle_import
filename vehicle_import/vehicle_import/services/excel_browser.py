import os
import jdatetime
from datetime import datetime
import openpyxl

import frappe
from frappe import _

EXCEL_FOLDER = os.path.join(
    frappe.get_site_path(),
    "private",
    "files",
    "excel_browser",
)

class ExcelBrowser:
    def get_files(self):
        if not os.path.isdir(EXCEL_FOLDER):
            return []

        files = []
        for filename in sorted(os.listdir(EXCEL_FOLDER)):

            if not filename.lower().endswith(".xlsx"):
                continue

            path = os.path.join(EXCEL_FOLDER, filename)

            stat = os.stat(path)

            modified = datetime.fromtimestamp(stat.st_mtime)

            if frappe.local.lang == "fa":
                modified = jdatetime.datetime.fromgregorian(
                    datetime=modified
                ).strftime("%Y/%m/%d %H:%M")
            else:
                modified = modified.strftime("%Y-%m-%d %H:%M")

            files.append({
                "file": filename,
                "size": stat.st_size,
                "modified": modified,
            })
        return files


    def get_table(self, filename):
        path = os.path.join(
            EXCEL_FOLDER,
            filename,
        )
        wb = openpyxl.load_workbook(
            path,
            data_only=True,
            read_only=True,
        )
        ws = wb.active
        headers = [
            str(cell.value or "")
            for cell in next(ws.iter_rows(max_row=1))
        ]
        columns = []
        for header in headers:
            columns.append({
                "id": header,
                "name": header,
            })        
        rows = []
        for values in ws.iter_rows(
                min_row=2,
                values_only=True):
            row = {}
            for header, value in zip(headers, values):
                row[header] = value
            rows.append(row)
        return {
            "columns": columns,
            "rows": rows,
        }


@frappe.whitelist()
def get_excel_files():
    return ExcelBrowser().get_files()


@frappe.whitelist()
def get_excel_table(filename):
    return ExcelBrowser().get_table(filename)