import frappe
from frappe import _
from frappe.query_builder import DocType
from pypika import Table, Case
from frappe.utils import fmt_money
import jdatetime

from io import BytesIO

import openpyxl
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment


class VehicleHolderCostReport:

    def execute(
        self,
        vehicle_holder,
    ):

        rows = self.get_rows(vehicle_holder)

        return {
            "columns": self.get_columns(),
            "rows": rows,
        }


    def get_columns(self):

        return [

            {
                "id": "vin",
                "name": _("VIN"),
            },

            {
                "id": "item",
                "name": _("Item"),
            },

            {
                "id": "status",
                "name": _("Status"),
            },

            {
                "id": "cost_date",
                "name": _("Cost Date"),
            },

            {
                "id": "cost_category",
                "name": _("Cost Category"),
            },

            {
                "id": "base_amount",
                "name": _("Amount"),
            },

            # {
            #     "id": "currency",
            #     "name": _("Foreign Currency"),
            #     "width": 90,
            # },

            # {
            #     "id": "foreign_amount",
            #     "name": _("Foreign Amount"),
            #     "width": 120,
            # },

            {
                "id": "cost_holder",
                "name": _("Document"),
            },

            {
                "id": "cost_reference_holder_type",
                "name": _("Type"),
            },

            {
                "id": "cost_holder_detail",
                "name": _("Detail Row"),
            },

            {
                "id": "creation_date",
                "name": _("Creation"),
            },


        ]


    def get_rows(
        self,
        vehicle_holder,
    ):

        CostLedger = DocType("Cost Ledger")
        CostEntry = DocType("Cost Entry")
        VehicleUnit = DocType("Vehicle Unit")
        VehicleHistory = DocType("Vehicle History")
        VehicleHolder = DocType("Vehicle Holder")
        VehicleHolderDetail = DocType("Vehicle Holder Detail")
        RefVehicleHolder = DocType("Vehicle Holder").as_("ref_vh")
        RefVehicleHolderDetail = DocType("Vehicle Holder Detail").as_("ref_vhd")
        RefVehicleHolderDetail_Holder = DocType("Vehicle Holder").as_("ref_vhd_h")

        rows = (

            frappe.qb

            .from_(VehicleHolder)

            .inner_join(VehicleHolderDetail)
            .on(
                VehicleHolderDetail.parent == VehicleHolder.name
            )

            .inner_join(VehicleHistory)
            .on(
                (VehicleHistory.parent == VehicleHolder.name)
                &
                (VehicleHistory.vehicle_history_vehicle_holder_detail == VehicleHolderDetail.name)
            )

            .inner_join(VehicleUnit)
            .on(
                VehicleUnit.name == VehicleHistory.vehicle_history_vehicle
            )

            .inner_join(CostLedger)
            .on(
                CostLedger.cost_ledger_vin == VehicleUnit.name
            )

            .inner_join(CostEntry)
            .on(
                CostEntry.name == CostLedger.cost_ledger_cost_entry
            )

            .left_join(RefVehicleHolder)
            .on(
                (CostEntry.cost_entry_reference_doctype == "Vehicle Holder")
                &
                (RefVehicleHolder.name == CostEntry.cost_entry_reference_name)
            )

            .left_join(RefVehicleHolderDetail)
            .on(
                (CostEntry.cost_entry_reference_doctype == "Vehicle Holder Detail")
                &
                (RefVehicleHolderDetail.name == CostEntry.cost_entry_reference_name)
            )

            .left_join(RefVehicleHolderDetail_Holder)
            .on(
                (CostEntry.cost_entry_reference_doctype == "Vehicle Holder Detail")
                &
                (RefVehicleHolderDetail_Holder.name == RefVehicleHolderDetail.parent)
            )

            .select(
                VehicleUnit.name.as_("vehicle"),
                Case()
                    .when(
                        (VehicleUnit.vehicle_vin.isnull())
                        |
                        (VehicleUnit.vehicle_vin == ""),
                        VehicleUnit.name,
                    )
                    .else_(VehicleUnit.vehicle_vin)
                .as_("vin"),
                VehicleHolderDetail.vehicle_holder_detail_item.as_("item"),
                CostEntry.docstatus.as_("status"),
                CostEntry.cost_entry_date.as_("cost_date"),
                CostEntry.creation.as_("creation_date"),
                CostEntry.cost_entry_cost_category.as_("cost_category"),
                CostLedger.cost_ledger_base_amount.as_("base_amount"),
                CostLedger.cost_ledger_currency.as_("currency"),
                CostLedger.cost_ledger_foreign_amount.as_("foreign_amount"),

                CostEntry.cost_entry_reference_name.as_("cost_reference"),
                CostEntry.cost_entry_reference_doctype.as_("cost_reference_doctype"),

                Case()
                    .when(
                        CostEntry.cost_entry_reference_doctype == "Vehicle Holder",
                        RefVehicleHolder.name,
                    )
                    .when(
                        CostEntry.cost_entry_reference_doctype == "Vehicle Holder Detail",
                        RefVehicleHolderDetail.parent,
                    )
                    .else_(None)
                .as_("cost_holder"),     

                Case()
                    .when(
                        CostEntry.cost_entry_reference_doctype == "Vehicle Holder",
                        None,
                    )
                    .when(
                        CostEntry.cost_entry_reference_doctype == "Vehicle Holder Detail",
                        CostEntry.cost_entry_reference_name,
                    )
                    .else_(None)
                .as_("cost_holder_detail"),     

                Case()
                    .when(
                        CostEntry.cost_entry_reference_doctype == "Vehicle Holder",
                        RefVehicleHolder.vehicle_holder_type,
                    )
                    .when(
                        CostEntry.cost_entry_reference_doctype == "Vehicle Holder Detail",
                        RefVehicleHolderDetail_Holder.vehicle_holder_type,
                    )
                    .else_(None)
                .as_("cost_reference_holder_type"),
            )

            .where(
                VehicleHolder.name == vehicle_holder
            )

            .orderby(
                CostLedger.creation
            )

            .run(
                as_dict=True,
            )
        )

        for row in rows:

            row["status"] = {
                0: _("Draft"),
                1: _("Submitted"),
                2: _("Cancelled"),
            }.get(row["status"])

            row["cost_category"] = _(row["cost_category"])

            row["base_amount_raw"] = row["base_amount"]
            row["base_amount"] = fmt_money(
                row["base_amount_raw"],
                precision=0,
            )

            row["cost_reference_holder_type"] = _(row["cost_reference_holder_type"])

            if frappe.local.lang == "fa":
                row["cost_date"] = jdatetime.date.fromgregorian(date=row["cost_date"]).strftime("%Y/%m/%d")
                row["creation_date"] = jdatetime.date.fromgregorian(date=row["creation_date"]).strftime("%Y/%m/%d")

        return rows


    def export_excel(
        self,
        vehicle_holder,
    ):

        rows = self.get_rows(vehicle_holder)
        columns = self.get_columns()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = _("Cost Report")

        #
        # Header
        #
        header_font = Font(
            bold=True,
        )

        header_alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        headers = ["#"] + [
            str(column["name"])
            for column in columns
        ]

        for col_index, title in enumerate(headers, start=1):

            cell = ws.cell(
                row=1,
                column=col_index,
                value=title,
            )

            cell.font = header_font
            cell.alignment = header_alignment

        #
        # Data
        #
        for row_index, row in enumerate(rows, start=2):

            ws.cell(
                row=row_index,
                column=1,
                value=row_index - 1,
            )

            for col_index, column in enumerate(columns, start=2):

                value = row.get(column["id"])

                if column["id"] == "base_amount":
                    value = row["base_amount_raw"]

                cell = ws.cell(
                    row=row_index,
                    column=col_index,
                    value=value,
                )

                if column["id"] == "base_amount":
                    cell.number_format = "#,##0"

        #
        # Table
        #        
        last_row = len(rows) + 1           # Header + Data
        last_col = len(columns) + 1        # + Row Number column
        table = Table(
            displayName="CostReport",
            ref=f"A1:{get_column_letter(last_col)}{last_row}",
        )
        style = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        table.tableStyleInfo = style
        table.headerRowCount = 1
        table.totalsRowShown = False

        ws.add_table(table)

        #
        # Total Row
        # 
        total_row = last_row + 1
        ws.cell(
            row=total_row,
            column=6,
            value=_("Total"),
        ).font = Font(bold=True)
        amount_col = 7    # با احتساب ستون شماره ردیف
        letter = get_column_letter(amount_col)
        cell = ws.cell(
            row=total_row,
            column=amount_col,
        )
        cell.value = f"=SUBTOTAL(9,{letter}2:{letter}{last_row})"
        cell.font = Font(bold=True)
        cell.number_format = '#,##0'

        #
        # Auto Filter
        #
        # ws.auto_filter.ref = ws.dimensions

        #
        # Freeze Header
        #
        ws.freeze_panes = "A2"

        #
        # Auto Width
        #
        for column_cells in ws.columns:

            length = max(
                len(str(cell.value or ""))
                for cell in column_cells
            )

            ws.column_dimensions[
                get_column_letter(column_cells[0].column)
            ].width = min(length + 3, 50)

        #
        # Download
        #
        output = BytesIO()
        wb.save(output)
        frappe.response.filename = (
            f"Cost Report - {vehicle_holder}.xlsx"
        )
        frappe.response.filecontent = output.getvalue()
        frappe.response.type = "binary"

@frappe.whitelist()
def get_vehicle_holder_cost_report(
    vehicle_holder,
):

    return VehicleHolderCostReport().execute(
        vehicle_holder,
    )


@frappe.whitelist()
def export_vehicle_holder_cost_report(
    vehicle_holder,
):

    VehicleHolderCostReport().export_excel(
        vehicle_holder,
    )