import frappe
from frappe import _
from frappe.query_builder import DocType
from pypika import Table, Case
from frappe.utils import fmt_money
import jdatetime

from io import BytesIO

import openpyxl
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment


class VehicleHolderCosts:

    def execute(
        self,
        vehicle_holder,
    ):

        rows = self.get_rows(vehicle_holder)

        return {
            "columns": self.get_columns(),
            "rows": rows,
        }


    def get_columns(self):

        return [

            {
                "id": "item",
                "name": _("Item"),
            },

            # {
            #     "id": "status",
            #     "name": _("Status"),
            # },

            # {
            #     "id": "cost_date",
            #     "name": _("Cost Date"),
            # },

            {
                "id": "cost_category",
                "name": _("Cost Category"),
            },

            {
                "id": "currency",
                "name": _("Currency"),
            },

            # {
            #     "id": "exchange_rate",
            #     "name": _("Rate"),
            # },

            {
                "id": "foreign_amount",
                "name": _("Foreign Amount"),
            },

            {
                "id": "base_amount",
                "name": _("Amount"),
            },

            # {
            #     "id": "creation_date",
            #     "name": _("Creation"),
            # },

            {
                "id": "cost_entry",
                "name": _("Cost Entry"),
            },

            {
                "id": "cost_holder_detail",
                "name": _("Detail Row"),
            },

            {
                "id": "description",
                "name": _("Description"),
            },

        ]


    def get_rows(
        self,
        vehicle_holder,
    ):

        CostEntry = DocType("Cost Entry")
        VehicleHolderDetail = DocType("Vehicle Holder Detail")

        #
        # Holder Costs
        #
        holder_rows = (
            frappe.qb
            .from_(CostEntry)
            .select(
                frappe.qb.terms.ValueWrapper("").as_("item"),
                CostEntry.docstatus.as_("status"),
                CostEntry.cost_entry_date.as_("cost_date"),
                CostEntry.creation.as_("creation_date"),
                CostEntry.cost_entry_cost_category.as_("cost_category"),
                CostEntry.cost_entry_currency.as_("currency"),
                CostEntry.cost_entry_exchange_rate.as_("exchange_rate"),
                CostEntry.cost_entry_foreign_amount.as_("foreign_amount"),
                CostEntry.cost_entry_base_amount.as_("base_amount"),
                CostEntry.cost_entry_description.as_("description"),                

                CostEntry.cost_entry_reference_name.as_("cost_reference"),
                CostEntry.cost_entry_reference_doctype.as_("cost_reference_doctype"),

                frappe.qb.terms.ValueWrapper(None).as_("cost_holder_detail"),

                CostEntry.name.as_("cost_entry"),
            )
            .where(
                (CostEntry.cost_entry_reference_doctype == "Vehicle Holder")
                &
                (CostEntry.cost_entry_reference_name == vehicle_holder)
            )
        ).run(as_dict=True)

        #
        # Detail Costs
        #
        detail_rows = (
            frappe.qb
            .from_(VehicleHolderDetail)
            .inner_join(CostEntry)
            .on(
                (CostEntry.cost_entry_reference_doctype == "Vehicle Holder Detail")
                &
                (CostEntry.cost_entry_reference_name == VehicleHolderDetail.name)
            )
            .select(
                VehicleHolderDetail.vehicle_holder_detail_item.as_("item"),
                CostEntry.docstatus.as_("status"),
                CostEntry.cost_entry_date.as_("cost_date"),
                CostEntry.creation.as_("creation_date"),
                CostEntry.cost_entry_cost_category.as_("cost_category"),
                CostEntry.cost_entry_currency.as_("currency"),
                CostEntry.cost_entry_exchange_rate.as_("exchange_rate"),
                CostEntry.cost_entry_foreign_amount.as_("foreign_amount"),
                CostEntry.cost_entry_base_amount.as_("base_amount"),

                CostEntry.cost_entry_reference_name.as_("cost_reference"),
                CostEntry.cost_entry_reference_doctype.as_("cost_reference_doctype"),

                CostEntry.cost_entry_reference_name.as_("cost_holder_detail"),

                CostEntry.name.as_("cost_entry"),
            )
            .where(
                VehicleHolderDetail.parent == vehicle_holder
            )
        ).run(as_dict=True)

        rows = holder_rows + detail_rows

        rows.sort(
            key=lambda r: (
                r["cost_date"],
                r["creation_date"],
            )
        )

        for row in rows:

            row["status"] = {
                0: _("Draft"),
                1: _("Submitted"),
                2: _("Cancelled"),
            }.get(row["status"])

            row["cost_category"] = _(row["cost_category"])

            row["exchange_rate_raw"] = row["exchange_rate"]
            row["exchange_rate"] = fmt_money(
                row["exchange_rate_raw"],
                precision=0,
            )

            row["foreign_amount_raw"] = row["foreign_amount"]
            row["foreign_amount"] = fmt_money(
                row["foreign_amount_raw"],
                precision=2,
            )

            row["base_amount_raw"] = row["base_amount"]
            row["base_amount"] = fmt_money(
                row["base_amount_raw"],
                precision=0,
            )

            if frappe.local.lang == "fa":
                row["cost_date"] = jdatetime.date.fromgregorian(
                    date=row["cost_date"]
                ).strftime("%Y/%m/%d")

                row["creation_date"] = jdatetime.date.fromgregorian(
                    date=row["creation_date"]
                ).strftime("%Y/%m/%d")

        return rows

    def export_excel(
        self,
        vehicle_holder,
    ):

        rows = self.get_rows(vehicle_holder)
        columns = self.get_columns()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = _("Costs")

        #
        # Header
        #
        header_font = Font(
            bold=True,
        )

        header_alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        headers = ["#"] + [
            str(column["name"])
            for column in columns
        ]

        for col_index, title in enumerate(headers, start=1):

            cell = ws.cell(
                row=1,
                column=col_index,
                value=title,
            )

            cell.font = header_font
            cell.alignment = header_alignment

        #
        # Data
        #
        for row_index, row in enumerate(rows, start=2):

            ws.cell(
                row=row_index,
                column=1,
                value=row_index - 1,
            )

            for col_index, column in enumerate(columns, start=2):

                value = row.get(column["id"])

                if column["id"] == "exchange_rate":
                    value = row["exchange_rate_raw"]

                if column["id"] == "foreign_amount":
                    value = row["foreign_amount_raw"]

                if column["id"] == "base_amount":
                    value = row["base_amount_raw"]

                cell = ws.cell(
                    row=row_index,
                    column=col_index,
                    value=value,
                )

                if column["id"] in ["exchange_rate", "foreign_amount", "base_amount"]:
                    cell.number_format = "#,##0"

        #
        # Table
        #        
        last_row = len(rows) + 1           # Header + Data
        last_col = len(columns) + 1        # + Row Number column
        table = Table(
            displayName="Costs",
            ref=f"A1:{get_column_letter(last_col)}{last_row}",
        )
        style = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        table.tableStyleInfo = style
        table.headerRowCount = 1
        table.totalsRowShown = False


        ws.add_table(table)

        #
        # Total Row
        # 
        total_row = last_row + 1
        amount_col = 6    # با احتساب ستون شماره ردیف
        ws.cell(
            row=total_row,
            column=amount_col-1,
            value=_("Total"),
        ).font = Font(bold=True)
        letter = get_column_letter(amount_col)
        cell = ws.cell(
            row=total_row,
            column=amount_col,
        )
        cell.value = f"=SUBTOTAL(9,{letter}2:{letter}{last_row})"
        cell.font = Font(bold=True)
        cell.number_format = '#,##0'

        #
        # Auto Filter
        #
        # ws.auto_filter.ref = ws.dimensions

        #
        # Freeze Header
        #
        ws.freeze_panes = "A2"

        #
        # Auto Width
        #
        for column_cells in ws.columns:

            length = max(
                len(str(cell.value or ""))
                for cell in column_cells
            )

            ws.column_dimensions[
                get_column_letter(column_cells[0].column)
            ].width = min(length + 3, 50)

        #
        # Download
        #
        output = BytesIO()
        wb.save(output)
        frappe.response.filename = (
            f"Costs - {vehicle_holder}.xlsx"
        )
        frappe.response.filecontent = output.getvalue()
        frappe.response.type = "binary"

@frappe.whitelist()
def get_vehicle_holder_costs(
    vehicle_holder,
):

    return VehicleHolderCosts().execute(
        vehicle_holder,
    )


@frappe.whitelist()
def export_vehicle_holder_costs(
    vehicle_holder,
):

    VehicleHolderCosts().export_excel(
        vehicle_holder,
    )